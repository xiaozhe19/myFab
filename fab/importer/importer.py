"""SMT2020 Excel 原始数据读取器。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from fab.importer.types import RawSheetData, RawSmtWorkbook


class Importer:
    """只读取 SMT 工作簿，不进行领域对象推导或单位换算。"""

    def __init__(self, file_path: Path) -> None:
        self.file_path = Path(file_path)

    def open_file(self) -> Workbook:
        """以流式只读模式打开一个 SMT ``.xlsx`` 工作簿。"""

        if not self.file_path.is_file():
            raise FileNotFoundError(f"找不到 SMT Excel 文件：{self.file_path}")
        if self.file_path.suffix.lower() != ".xlsx":
            raise ValueError(f"只支持 .xlsx 文件：{self.file_path.name}")
        return load_workbook(self.file_path, read_only=True, data_only=True)

    def read(self) -> RawSmtWorkbook:
        """读取所有工作表的原始表头和非空数据行。"""

        workbook = self.open_file()
        try:
            return RawSmtWorkbook(
                source_path=self.file_path,
                sheets=tuple(self.read_sheet(workbook[name]) for name in workbook.sheetnames),
            )
        finally:
            workbook.close()

    def read_sheet(self, sheet: Worksheet) -> RawSheetData:
        """读取一张工作表，不解释列含义，也不修改单元格值。"""

        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = next(iterator)
        except StopIteration:
            raise ValueError(f"工作表为空：{sheet.title}") from None
        rows = tuple(
            tuple(values)
            for values in iterator
            if values and any(value is not None for value in values)
        )
        return RawSheetData(
            sheet.title,
            tuple("" if value is None else str(value) for value in headers),
            rows,
        )
